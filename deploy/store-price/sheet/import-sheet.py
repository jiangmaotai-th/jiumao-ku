#!/usr/bin/env python3
"""Convert AI product tier xlsx → latest.raw.json for apply-sheet.mjs."""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print('需要 openpyxl：pip3 install openpyxl', file=sys.stderr)
    sys.exit(1)

HERE = Path(__file__).resolve().parent
KEYS = [
    'category',
    'name',
    'planStructure',
    'freeTrial',
    'personalPlans',
    'teamPlans',
    'apiPlans',
    'buyout',
    'statusNote',
    'changeNote',
]


def main() -> int:
    ap = argparse.ArgumentParser(description='Import AI tier spreadsheet')
    ap.add_argument(
        'xlsx',
        nargs='?',
        default=str(HERE / 'latest.xlsx'),
        help='xlsx path (default: sheet/latest.xlsx)',
    )
    ap.add_argument(
        '-o',
        '--out',
        default=str(HERE / 'latest.raw.json'),
        help='output JSON path',
    )
    args = ap.parse_args()
    src = Path(args.xlsx).expanduser().resolve()
    if not src.exists():
        print(f'找不到表格: {src}', file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[wb.sheetnames[0]]
    products = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 5:
            continue
        vals = [(c.strip() if isinstance(c, str) else c) for c in row]
        if not vals or not vals[1]:
            continue
        obj = {}
        for k, v in zip(KEYS, vals[:10]):
            obj[k] = '' if v is None else str(v).strip()
        products.append(obj)

    out = {
        'updated': date.today().isoformat(),
        'source': src.name,
        'products': products,
    }
    out_path = Path(args.out)
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    # Keep a copy as latest.xlsx when importing from elsewhere
    if src.resolve() != (HERE / 'latest.xlsx').resolve():
        import shutil

        shutil.copy2(src, HERE / 'latest.xlsx')
    print(f'OK {len(products)} products → {out_path}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
